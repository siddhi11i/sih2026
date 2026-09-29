import io
import csv
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_csv_export(standards: List[Dict[str, Any]], query: str = "") -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Procurement Query", query])
    writer.writerow([])
    writer.writerow(["Rank", "IS Number", "Title", "Year", "Sector / Category", "Match Percentage", "Relevance Tier", "Data Status", "Why Matched"])
    for idx, s in enumerate(standards, 1):
        writer.writerow([
            idx,
            s.get("is_number", ""),
            s.get("title", ""),
            s.get("year", ""),
            s.get("category", ""),
            f"{s.get('match_pct', 0)}%",
            s.get("match_tier", ""),
            s.get("data_status", "REAL"),
            "; ".join(s.get("why_matched", []))
        ])
    return output.getvalue()

def generate_docx_tender_report(
    standards: List[Dict[str, Any]],
    query: str,
    organization: str = "National Public Procurement Authority",
    tender_ref: str = "TND-BIS-2026-F&D"
) -> bytes:
    doc = Document()

    # Title & Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("BUREAU OF INDIAN STANDARDS (BIS)\nTECHNICAL STANDARDS CONFORMANCE ANNEXURE")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Government of India • Department of Consumer Affairs • Food & Dairy Division\n")
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_heading("1. Tender Scope & Technical Procurement Statement", level=1)
    doc.add_paragraph(f"Organization / Procuring Body: {organization}")
    doc.add_paragraph(f"Tender Reference: {tender_ref}")
    doc.add_paragraph(f"Procurement Scope Statement: \"{query}\"")

    doc.add_heading("2. Mandatory & Applicable Indian Standards (IS Annexure)", level=1)
    doc.add_paragraph(
        "The following Indian Standards are retrieved from the official BIS repository and are recommended for "
        "incorporation into the technical specifications and evaluation criteria of the tender documentation:"
    )

    # Table of Standards
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Sr.", "IS Number", "Standard Title & Revision", "Category", "Match %", "Data Governance Status"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9)

    for idx, s in enumerate(standards, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = str(s.get("is_number", ""))
        row_cells[2].text = f"{s.get('title', '')} ({s.get('year', 'N/A')})"
        row_cells[3].text = str(s.get("category", "Food & Dairy"))
        row_cells[4].text = f"{s.get('match_pct', 0)}% [{s.get('match_tier', '')}]"
        
        status_text = str(s.get("data_status", "REAL"))
        row_cells[5].text = status_text
        if "SAMPLE" in status_text or "UNVERIFIED" in status_text:
            row_cells[5].paragraphs[0].runs[0].font.color.rgb = RGBColor(180, 83, 9)

    doc.add_heading("3. Regulatory & Quality Conformance Notice", level=1)
    notice_p = doc.add_paragraph(
        "IMPORTANT NOTICE TO TENDER COMMITTEES:\n"
        "1. All product standards referenced above must be verified for the latest amendments and active Gazette notifications.\n"
        "2. Any entry tagged 'SAMPLE - verify before use' represents a schema template and must undergo committee verification prior to tender gazette publication.\n"
        "3. Mandatory Clause 2 allied test methods and sampling procedures (IS 11721, IS 11546) must be adhered to during pre-dispatch quality inspections."
    )
    notice_p.runs[0].font.size = Pt(9.5)
    notice_p.runs[0].font.italic = True

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
